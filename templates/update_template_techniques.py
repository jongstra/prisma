"""
Refresh the technique list in the MaGMa Excel template from PRISMA's ATT&CK catalog.

Run this after a new tactics_and_techniques_by_domain.json has been generated (see tools/build_attack_catalog.ipynb):

    tools/.venv/bin/python templates/update_template_techniques.py 17.1

It rewrites the 'Mitre Techniques' sheet: every enterprise technique and sub-technique, with its tactics in Unified
Kill Chain order (the order of the 'ATTACK and Kill Chain mapping' sheet). The L3 sheet uses the first of these tactics
to suggest an L2 use case. When LibreOffice is installed, the workbook is recalculated, so the values that the
converter notebook reads are up to date; otherwise, open and save the file once in Excel or LibreOffice.
"""
import json, shutil, subprocess, sys, tempfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'templates' / 'magma_template.xlsx'
CATALOG = ROOT / 'frontend' / 'public' / 'tactics_and_techniques_by_domain.json'
LIBREOFFICE = shutil.which('soffice') or '/Applications/LibreOffice.app/Contents/MacOS/soffice'


def techniques_with_tactics(catalog_path, tactic_rank):
    """Return (ID, name, tactics) for every enterprise (sub-)technique, sorted by ID, tactics in kill chain order."""
    tactics_of, names = {}, {}
    for tactic in json.loads(catalog_path.read_text())['enterprise']['tactics']:
        for technique in tactic['techniques']:
            items = [(technique, technique['name'])]
            items += [(sub, f"{technique['name']}: {sub['name']}") for sub in technique.get('sub_techniques') or []]
            for item, name in items:
                tactics_of.setdefault(item['external_id'], set()).add(tactic['name'])
                names[item['external_id']] = name
    unknown = {tactic for tactics in tactics_of.values() for tactic in tactics} - set(tactic_rank)
    if unknown:
        raise ValueError(f"Add these tactics to the 'ATTACK and Kill Chain mapping' sheet first: {sorted(unknown)}")
    order = lambda technique_id: [int(part) for part in technique_id[1:].split('.')]
    return [(i, names[i], ', '.join(sorted(tactics_of[i], key=tactic_rank.get))) for i in sorted(tactics_of, key=order)]


def main(attack_version):
    workbook = openpyxl.load_workbook(TEMPLATE)

    # The kill chain order of the tactics, from the mapping sheet (one row per Unified Kill Chain step).
    mapping = workbook['ATTACK and Kill Chain mapping'].iter_rows(min_row=2, max_col=1, values_only=True)
    tactic_rank = {tactic: rank for rank, (tactic,) in enumerate(mapping) if tactic}
    rows = techniques_with_tactics(CATALOG, tactic_rank)

    # Replace the sheet (at the same position), so no old rows are left behind.
    index = workbook.sheetnames.index('Mitre Techniques')
    workbook.remove(workbook['Mitre Techniques'])
    sheet = workbook.create_sheet('Mitre Techniques', index)
    sheet.append(['ID', 'Name', 'Tactics (Unified Kill Chain order)'])
    for row in rows:
        sheet.append(row)
    table = Table(displayName='Techniques', ref=f'A1:C{len(rows) + 1}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    sheet['E1'] = f'MITRE ATT&CK Enterprise v{attack_version}, from PRISMA (frontend/public/tactics_and_techniques_by_domain.json)'
    sheet['E1'].font = Font(italic=True)
    for column, width in zip('ABC', (12, 70, 60)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'

    with tempfile.TemporaryDirectory() as folder:
        saved = Path(folder) / TEMPLATE.name
        workbook.save(saved)
        if Path(LIBREOFFICE).exists():
            # Recalculate and save with LibreOffice; openpyxl does not store the results of formulas.
            out = Path(folder) / 'recalculated'
            profile = Path(folder) / 'profile'
            (profile / 'user').mkdir(parents=True)
            (profile / 'user' / 'registrymodifications.xcu').write_text(
                '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item></oor:items>')
            subprocess.run([LIBREOFFICE, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--calc',
                            '--convert-to', 'xlsx:Calc Office Open XML', '--outdir', str(out), str(saved)],
                           check=True, capture_output=True)
            saved = out / TEMPLATE.name
        else:
            print('LibreOffice was not found: open and save the template once in Excel or LibreOffice.')
        shutil.copy(saved, TEMPLATE)
    print(f'{TEMPLATE.name}: {len(rows)} techniques and sub-techniques (ATT&CK v{attack_version}).')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit('Usage: update_template_techniques.py <ATT&CK version, e.g. 17.1>')
    main(sys.argv[1])
